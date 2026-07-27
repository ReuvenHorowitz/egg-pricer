"""Build the formatted Excel review workbook."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.config import (
    CURRENCY_FORMAT,
    ITEM_MASTER_PATH,
    PERCENT_FORMAT,
    PRICE_HISTORY_PATH,
    RED_FILL,
    ROUNDING_INCREMENT,
    TRACKED_VENDOR_NUMBERS,
    YELLOW_FILL,
)
from src.parse_martins_pdf import PdfMetadata, ParsedPdfRow
from src.price_math import PricedRow


REVIEW_HEADERS = [
    "Item Number",
    "Description",
    "Vendor Number",
    "Raw List Price",
    "Cost Multiplier",
    "Current Cost",
    "New Cost",
    "Cost Change $",
    "Cost Change %",
    "Current Price",
    "Suggested Price",
    "Action",
    "Notes",
]

RAW_HEADERS = [
    "Vendor Number",
    "Raw List Price",
    "Description From PDF",
    "Page",
    "Matched Line",
    "Matched",
]


def _style_header(sheet, row_number: int = 1) -> None:
    for cell in sheet[row_number]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"


def _autosize_columns(sheet, max_width: int = 45) -> None:
    for column_cells in sheet.columns:
        letter = get_column_letter(column_cells[0].column)
        width = max(len(str(cell.value or "")) for cell in column_cells) + 2
        sheet.column_dimensions[letter].width = min(width, max_width)


def _set_currency(cell) -> None:
    cell.number_format = CURRENCY_FORMAT


def _set_percent(cell) -> None:
    cell.number_format = PERCENT_FORMAT


def _build_review_sheet(sheet, priced_rows: list[PricedRow]) -> None:
    sheet.title = "Price Change Review"
    sheet.append(REVIEW_HEADERS)
    _style_header(sheet)

    yellow_fill = PatternFill(fill_type="solid", fgColor=YELLOW_FILL)
    red_fill = PatternFill(fill_type="solid", fgColor=RED_FILL)

    for priced in priced_rows:
        item = priced.item
        row = [
            item.item_number,
            item.description,
            item.vendor_number,
            priced.raw_list_price,
            item.cost_multiplier,
            item.current_cost,
            priced.new_cost,
            priced.cost_change_dollar,
            priced.cost_change_pct,
            priced.current_price,
            priced.suggested_price,
            priced.action,
            priced.notes,
        ]
        sheet.append(row)
        row_idx = sheet.max_row

        for col_idx in (4, 6, 7, 8, 10, 11):
            _set_currency(sheet.cell(row=row_idx, column=col_idx))
        _set_percent(sheet.cell(row=row_idx, column=9))

        missing = priced.raw_list_price is None
        changed = (
            priced.new_cost is not None
            and item.current_cost is not None
            and abs(priced.new_cost - item.current_cost) > 0.001
        )

        if missing:
            for col_idx in range(1, len(REVIEW_HEADERS) + 1):
                sheet.cell(row=row_idx, column=col_idx).fill = red_fill
        elif changed:
            sheet.cell(row=row_idx, column=7).fill = yellow_fill

    _autosize_columns(sheet)


def _build_raw_extract_sheet(sheet, parsed_rows: dict[str, ParsedPdfRow]) -> None:
    sheet.title = "Raw Extract"
    sheet.append(RAW_HEADERS)
    _style_header(sheet)

    for vendor in TRACKED_VENDOR_NUMBERS:
        parsed = parsed_rows.get(vendor.upper())
        if parsed is None:
            sheet.append([vendor, None, "", "", "", "No"])
            continue

        sheet.append(
            [
                parsed.vendor_number,
                parsed.raw_list_price,
                parsed.description_from_pdf,
                parsed.source_page or "",
                parsed.source_line,
                "Yes" if parsed.matched else "No",
            ]
        )
        _set_currency(sheet.cell(row=sheet.max_row, column=2))

    _autosize_columns(sheet)


def _build_settings_sheet(
    sheet,
    pdf_metadata: PdfMetadata,
    pdf_path: Path,
    output_path: Path,
    priced_rows: list[PricedRow],
) -> None:
    sheet.title = "Settings"
    run_time = datetime.now()

    settings_rows = [
        ("Run Timestamp", run_time.isoformat(timespec="seconds")),
        ("Run Date", run_time.date().isoformat()),
        ("PDF Filename", pdf_metadata.filename),
        ("PDF Path", str(pdf_path)),
        ("PDF Customer Code", pdf_metadata.customer_code or ""),
        ("PDF Price Date", pdf_metadata.price_date or ""),
        ("Item Master Path", str(ITEM_MASTER_PATH)),
        ("Item Master Rows Loaded", len(priced_rows)),
        ("Tracked Vendor Numbers", ", ".join(TRACKED_VENDOR_NUMBERS)),
        ("Rounding Rule", f"Nearest ${ROUNDING_INCREMENT:.2f}"),
        ("Price History Path", str(PRICE_HISTORY_PATH)),
        ("Output Workbook", str(output_path)),
        ("Pricing Formula", "new_cost = raw_list_price * cost_multiplier"),
        (
            "Suggested Price Formula",
            "Preserve current margin when current_price and current_cost exist; otherwise use target_margin",
        ),
        ("Business Central Update", "v1 review only — no automatic BC update"),
    ]

    sheet.append(["Setting", "Value"])
    _style_header(sheet)
    for setting, value in settings_rows:
        sheet.append([setting, value])

    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 80


def build_excel_report(
    priced_rows: list[PricedRow],
    parsed_rows: dict[str, ParsedPdfRow],
    pdf_metadata: PdfMetadata,
    pdf_path: Path,
    output_path: Path,
) -> Path:
    """Create the three-sheet Excel review workbook."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    review_sheet = workbook.active
    raw_sheet = workbook.create_sheet("Raw Extract")
    settings_sheet = workbook.create_sheet("Settings")

    _build_review_sheet(review_sheet, priced_rows)
    _build_raw_extract_sheet(raw_sheet, parsed_rows)
    _build_settings_sheet(settings_sheet, pdf_metadata, pdf_path, output_path, priced_rows)

    workbook.save(output_path)
    return output_path
