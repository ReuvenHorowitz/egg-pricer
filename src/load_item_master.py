"""Load and validate the item master spreadsheet."""
from __future__ import annotations

from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook

from src.config import ITEM_MASTER_COLUMNS, ITEM_MASTER_PATH, TRACKED_VENDOR_NUMBERS


class ItemMasterError(Exception):
    """Raised when item master data is missing or invalid."""


@dataclass(frozen=True)
class ItemMasterRow:
    item_number: str
    description: str
    vendor_number: str
    pack_size: str | None
    cost_multiplier: float
    current_cost: float | None
    current_price: float | None
    target_margin: float
    rounding_rule: str | None
    source_row: int


def _normalize_header(value: Any) -> str:
    return str(value or "").strip().lower()


def _parse_float(value: Any, field_name: str, row_number: int, required: bool = True) -> float | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        if required:
            raise ItemMasterError(f"Row {row_number}: missing required value for '{field_name}'")
        return None
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ItemMasterError(f"Row {row_number}: invalid number for '{field_name}': {value!r}") from exc


def _parse_text(value: Any, field_name: str, row_number: int, required: bool = True) -> str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        if required:
            raise ItemMasterError(f"Row {row_number}: missing required value for '{field_name}'")
        return None
    return str(value).strip()


def load_item_master() -> list[ItemMasterRow]:
    """Load tracked rows from the item master workbook."""
    if not ITEM_MASTER_PATH.exists():
        raise ItemMasterError(f"Item master not found: {ITEM_MASTER_PATH}")

    try:
        workbook = load_workbook(ITEM_MASTER_PATH, read_only=True, data_only=True)
    except PermissionError as exc:
        raise ItemMasterError(
            f"Cannot read item master (file may be open in Excel): {ITEM_MASTER_PATH}. "
            "Close the workbook and run again."
        ) from exc
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    if not rows:
        raise ItemMasterError("Item master workbook is empty")

    header = [_normalize_header(cell) for cell in rows[0]]
    missing_columns = [column for column in ITEM_MASTER_COLUMNS if column not in header]
    if missing_columns:
        raise ItemMasterError(f"Item master missing columns: {', '.join(missing_columns)}")

    column_index = {name: header.index(name) for name in ITEM_MASTER_COLUMNS}
    tracked = {vendor.upper() for vendor in TRACKED_VENDOR_NUMBERS}
    loaded: list[ItemMasterRow] = []

    for row_number, raw_row in enumerate(rows[1:], start=2):
        if not raw_row or all(cell is None or str(cell).strip() == "" for cell in raw_row):
            continue

        vendor_number = _parse_text(raw_row[column_index["vendor_number"]], "vendor_number", row_number)
        if vendor_number is None or vendor_number.upper() not in tracked:
            continue

        item = ItemMasterRow(
            item_number=_parse_text(raw_row[column_index["item_number"]], "item_number", row_number) or "",
            description=_parse_text(raw_row[column_index["description"]], "description", row_number) or "",
            vendor_number=vendor_number.upper(),
            pack_size=_parse_text(raw_row[column_index["pack_size"]], "pack_size", row_number, required=False),
            cost_multiplier=_parse_float(
                raw_row[column_index["cost_multiplier"]], "cost_multiplier", row_number
            ),
            current_cost=_parse_float(raw_row[column_index["current_cost"]], "current_cost", row_number, required=False),
            current_price=_parse_float(
                raw_row[column_index["current_price"]], "current_price", row_number, required=False
            ),
            target_margin=_parse_float(raw_row[column_index["target_margin"]], "target_margin", row_number),
            rounding_rule=_parse_text(
                raw_row[column_index["rounding_rule"]], "rounding_rule", row_number, required=False
            ),
            source_row=row_number,
        )
        loaded.append(item)

    found_vendors = {item.vendor_number for item in loaded}
    for vendor in TRACKED_VENDOR_NUMBERS:
        if vendor.upper() not in found_vendors:
            print(f"WARNING: Tracked vendor {vendor} not found in item master")

    if not loaded:
        raise ItemMasterError("No tracked vendor rows found in item master")

    order = {vendor: index for index, vendor in enumerate(TRACKED_VENDOR_NUMBERS)}
    loaded.sort(key=lambda item: order.get(item.vendor_number, 999))
    return loaded
